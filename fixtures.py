"""Fixture workbooks for the auditor tests: one deliberately broken, one clean, one with an external link."""
import os, sys
from openpyxl import Workbook


def broken(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Price", "Total", "Tax", "Lookup", "Misc"])
    items = ["A", "B", "C", "D", "E", "F", "G", "H"]
    for i, it in enumerate(items, start=2):
        ws.cell(i, 1, it); ws.cell(i, 2, i); ws.cell(i, 3, 10 + i)
        ws.cell(i, 4, f"=B{i}*C{i}")
        ws.cell(i, 5, f"=D{i}*0.0825")
    ws["D5"] = 480                      # typed over a formula
    ws["E6"] = "=D6*0.09"               # pattern break between agreeing neighbours
    ws["F2"] = "=VLOOKUP(A2,Prices!A1:B3,2)"            # approximate lookup
    ws["F3"] = "=VLOOKUP(A3,Prices!A1:B3,2,FALSE)"      # fine
    ws["G2"] = "=C2/(B2-B2)"            # #DIV/0!
    ws["G3"] = '=INDIRECT("A"&2)'       # volatile
    p = wb.create_sheet("Prices")
    for r in (("A", 1), ("B", 2), ("C", 3)):
        p.append(list(r))
    h = wb.create_sheet("Scratch")
    h["A1"] = 1
    h.sheet_state = "hidden"
    wb.save(path)


def clean(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Tax rate", 0.0825
    ws.append(["Item", "Qty", "Price", "Total", "Tax"])
    for i in range(3, 9):
        ws.cell(i, 1, f"Item {i}"); ws.cell(i, 2, i); ws.cell(i, 3, 10 + i)
        ws.cell(i, 4, f"=B{i}*C{i}")
        ws.cell(i, 5, f"=D{i}*$B$1")
    ws["D9"] = "=SUM(D3:D8)"
    wb.save(path)


def extlink(path):
    wb = Workbook()
    wb.active["A1"] = "='C:\\Data\\[budget.xlsx]Sheet1'!A1"
    wb.save(path)


def write_all(folder):
    os.makedirs(folder, exist_ok=True)
    for name, fn in (("broken.xlsx", broken), ("clean.xlsx", clean), ("ext.xlsx", extlink)):
        fn(os.path.join(folder, name))


if __name__ == "__main__":
    write_all(sys.argv[1])
    print("written", sys.argv[1])
