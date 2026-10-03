#!/usr/bin/env python3
"""xlsx_audit.py - static + recalculation audit of an .xlsx workbook. Local only, read-only, no network.

usage: python xlsx_audit.py <workbook.xlsx> [--json out.json] [--no-recalc] [--max-cells 200000]

Checks: error values after an independent recalculation, hidden sheets, rows and columns, external links,
volatile functions, formulas with hard-coded constants, formulas that break the pattern of their neighbours,
typed-over formulas inside a formula column, approximate-match lookups, broken defined names.
Heuristic: every finding needs a human look. It cannot prove a model is right, only point at where models usually go wrong.
"""
import argparse, json, os, re, shutil, sys, tempfile, warnings
from collections import Counter, defaultdict

warnings.filterwarnings("ignore")
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

SEV = {"HIGH": 0, "MEDIUM": 1, "LOW": 2, "INFO": 3}
VOLATILE = ("INDIRECT", "OFFSET", "TODAY", "NOW", "RAND", "RANDBETWEEN", "INFO", "CELL")
ERR_RE = re.compile(r"^#(REF!|DIV/0!|N/A|VALUE!|NAME\?|NUM!|NULL!|SPILL!|CALC!)")
OK_CONSTS = {"0", "1", "2", "3", "4", "5", "7", "10", "12", "52", "100", "1000", "365", "24", "60", "0.5", "-1"}
STR_RE = re.compile(r'"[^"]*"')
REF_RE = re.compile(r"(?:'[^']+'|[A-Za-z0-9_]+)?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
NUM_RE = re.compile(r"(?<![A-Za-z0-9_.$!:'])-?\d+(?:\.\d+)?(?![A-Za-z0-9_(!:])")


class Finding(dict):
    def __init__(self, rule, severity, title, sheet, cell, detail, fix):
        super().__init__(rule=rule, severity=severity, title=title, sheet=sheet, cell=cell, detail=detail, fix=fix)


TOL_RE = re.compile(r"(?:<=|>=|<|>|=)\s*-?0?\.00\d+")


def magic_numbers(formula):
    f = STR_RE.sub('""', formula)
    f = TOL_RE.sub("TOL", f)  # rounding tolerances such as <=0.005 are not business assumptions
    f = REF_RE.sub("REF", f)
    return [n for n in NUM_RE.findall(f) if n.lstrip("-") not in OK_CONSTS and n not in OK_CONSTS]


def same_pattern(a, b):
    """True if formula cell b is exactly what filling cell a down/across to b would produce."""
    try:
        return Translator(a.value, origin=a.coordinate).translate_formula(b.coordinate) == b.value
    except Exception:
        return False


def is_f(c):
    return c is not None and isinstance(c.value, str) and c.value.startswith("=")


def recalc_errors(path):
    """Independent recalculation with the pure-python `formulas` engine; returns list of (sheet, cell, error)."""
    try:
        import formulas
    except Exception:
        return None
    out = tempfile.mkdtemp()
    try:
        xl = formulas.ExcelModel().loads(path).finish()
        xl.calculate()
        xl.write(dirpath=out)
        files = [f for f in os.listdir(out) if f.lower().endswith(".xlsx")]
        wb = load_workbook(os.path.join(out, files[0]), data_only=True)
        bad = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is not None and (ERR_RE.match(str(v)) or "XlError" in type(v).__name__):
                        bad.append((ws.title, c.coordinate, str(v)))
        return bad
    except Exception as e:  # engine could not handle the workbook
        return f"recalculation skipped: {type(e).__name__}: {str(e)[:120]}"
    finally:
        shutil.rmtree(out, ignore_errors=True)


def audit(path, recalc=True, max_cells=200000):
    F = []
    wb = load_workbook(path)  # formulas as written
    try:
        wbv = load_workbook(path, data_only=True)  # cached values (present if last saved by Excel/Sheets)
    except Exception:
        wbv = None
    n_cells = n_formulas = 0

    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            F.append(Finding("HIDDEN-SHEET", "MEDIUM", f"Sheet is {ws.sheet_state}", ws.title, "", "Hidden sheets can hold inputs or logic nobody reviews.", "Unhide, review, document or delete."))
        hid_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
        hid_rows = [k for k, d in ws.row_dimensions.items() if d.hidden]
        if hid_cols or hid_rows:
            F.append(Finding("HIDDEN-ROWCOL", "LOW", "Hidden rows or columns", ws.title, "", f"columns {hid_cols[:8]} rows {hid_rows[:8]}", "Check hidden cells for logic or stale data."))

        grid = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                n_cells += 1
                if n_cells > max_cells:
                    break
                grid[(c.row, c.column)] = c
            if n_cells > max_cells:
                F.append(Finding("TRUNCATED", "INFO", "Cell limit reached", ws.title, "", f"audited the first {max_cells} cells only", "Raise --max-cells."))
                break

        for (r, col), c in grid.items():
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                n_formulas += 1
                up = v.upper()
                if re.search(r"\[\d+\]|'?[A-Za-z]:\\|https?://", v):
                    F.append(Finding("EXT-LINK", "HIGH", "Formula links to another file", ws.title, c.coordinate, v[:100], "Replace the link with values or bring the data in; links break when files move."))
                for fn in VOLATILE:
                    if re.search(rf"\b{fn}\(", up):
                        F.append(Finding("VOLATILE", "LOW", f"Volatile function {fn}", ws.title, c.coordinate, v[:100], "Volatile functions recalc constantly and make results hard to audit or reproduce."))
                        break
                m = magic_numbers(v)
                if m:
                    F.append(Finding("MAGIC-NUMBER", "MEDIUM", "Hard-coded number inside a formula", ws.title, c.coordinate, f"{v[:90]}  (numbers: {', '.join(m[:4])})", "Move assumptions like rates and thresholds into labelled input cells."))
                if re.search(r"VLOOKUP\([^()]*(?:\([^()]*\)[^()]*)*\)", up):
                    mm = re.search(r"VLOOKUP\(((?:[^()]|\([^()]*\))*)\)", up)
                    if mm:
                        args = [a for a in re.split(r",(?![^()]*\))", mm.group(1))]
                        if len(args) < 4 or args[3].strip() in ("TRUE", "1"):
                            F.append(Finding("LOOKUP-APPROX", "MEDIUM", "VLOOKUP with approximate match", ws.title, c.coordinate, v[:100], "Use FALSE (exact match) unless the table is sorted on purpose; approximate match returns wrong rows silently."))

        # pattern breaks: the odd one out in the middle of a filled-down run (2 consistent cells above, 2 below)
        for (r, col), c in grid.items():
            if not is_f(c):
                continue
            u2, u1, d1, d2 = (grid.get((r + k, col)) for k in (-2, -1, 1, 2))
            if all(is_f(x) for x in (u2, u1, d1, d2)) and same_pattern(u2, u1) and same_pattern(d1, d2) and same_pattern(u1, d1) and not same_pattern(u1, c):
                try:
                    expect = Translator(u1.value, origin=u1.coordinate).translate_formula(c.coordinate)
                except Exception:
                    continue
                F.append(Finding("PATTERN-BREAK", "HIGH", "Formula differs from its neighbours", ws.title, c.coordinate, f"has {c.value[:70]} but neighbours imply {expect[:70]}", "Check whether this is intentional; a lone different formula in a filled-down column is a classic source of wrong totals."))

        # typed-over formulas: a constant number where the neighbours' pattern says a formula belongs
        for (r, col), c in grid.items():
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                u2, u1, d1, d2 = (grid.get((r + k, col)) for k in (-2, -1, 1, 2))
                if all(is_f(x) for x in (u1, d1)) and same_pattern(u1, d1) and (not is_f(u2) or same_pattern(u2, u1)) and (not is_f(d2) or same_pattern(d1, d2)):
                    F.append(Finding("TYPED-OVER", "HIGH", "Hard-coded number inside a filled-down formula run", ws.title, c.coordinate, f"value {v} sits between {u1.coordinate} and {d1.coordinate}, which follow one pattern", "Likely a formula overwritten with a typed value; restore the formula."))

        # cached error values (only if the file was last saved by an app that stores them)
        if wbv is not None and ws.title in wbv.sheetnames:
            for row in wbv[ws.title].iter_rows():
                for c in row:
                    if c.value is not None and ERR_RE.match(str(c.value)):
                        F.append(Finding("CACHED-ERROR", "HIGH", f"Error value {c.value}", ws.title, c.coordinate, "stored result of last save", "Fix the formula or its inputs."))

    # defined names
    try:
        names = wb.defined_names
        items = names.items() if hasattr(names, "items") else [(n.name, n) for n in names.definedName]
        for nm, dn in items:
            if "#REF!" in str(dn.attr_text):
                F.append(Finding("NAME-REF", "HIGH", f"Defined name '{nm}' points to #REF!", "", "", str(dn.attr_text)[:100], "Delete or repoint the name."))
    except Exception:
        pass

    note = None
    if recalc:
        res = recalc_errors(path)
        if isinstance(res, list):
            for sh, cell, err in res:
                F.append(Finding("CALC-ERROR", "HIGH", f"Error after recalculation: {err}", sh, cell, "independent recalculation with the `formulas` engine", "Fix the formula or its inputs."))
        elif isinstance(res, str):
            note = res
        else:
            note = "recalculation skipped: `formulas` package not installed"
    seen, uniq = set(), []
    for f in F:
        k = (f["rule"], f["sheet"], f["cell"])
        if k not in seen:
            seen.add(k)
            uniq.append(f)
    uniq.sort(key=lambda f: (SEV[f["severity"]], f["sheet"], f["cell"]))
    grouped, index = [], {}
    for f in uniq:  # one entry per (rule, sheet): first cell plus a count, so 700 repeats do not bury the real problems
        k = (f["rule"], f["sheet"], f["severity"])
        if k in index and f["rule"] in ("MAGIC-NUMBER", "VOLATILE", "LOOKUP-APPROX"):
            g = grouped[index[k]]
            g["count"] += 1
            if len(g["other_cells"]) < 8:
                g["other_cells"].append(f["cell"])
        else:
            index[k] = len(grouped)
            grouped.append(dict(f, count=1, other_cells=[]))
    uniq = grouped
    return dict(workbook=os.path.abspath(path), sheets=[w.title for w in wb.worksheets], cells=n_cells, formulas=n_formulas,
                counts=dict(Counter({sv: sum(x["count"] for x in uniq if x["severity"] == sv) for sv in SEV if any(x["severity"] == sv for x in uniq)})), recalculation_note=note, findings=uniq)


def inspect_cell(path, sheet, cell):
    """Formula, cached value and the cells it reads, for one cell."""
    wb = load_workbook(path)
    if sheet not in wb.sheetnames:
        raise ValueError(f"no sheet named {sheet!r}; sheets: {wb.sheetnames}")
    c = wb[sheet][cell]
    cached = None
    try:
        cached = load_workbook(path, data_only=True)[sheet][cell].value
    except Exception:
        pass
    v = c.value
    refs = []
    if isinstance(v, str) and v.startswith("="):
        refs = sorted(set(m.group(0) for m in REF_RE.finditer(STR_RE.sub('""', v)) if re.search(r"\d", m.group(0))))
    return dict(sheet=sheet, cell=cell, value_or_formula=v, is_formula=isinstance(v, str) and v.startswith("="), cached_value=cached, reads=refs)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--cell", help="Sheet!A1: print one cell's formula, cached value and the cells it reads, then exit")
    ap.add_argument("--json")
    ap.add_argument("--no-recalc", action="store_true")
    ap.add_argument("--max-cells", type=int, default=200000)
    a = ap.parse_args()
    if not os.path.isfile(a.path):
        sys.exit(f"not a file: {a.path}")
    if a.cell:
        sh, _, ce = a.cell.rpartition("!")
        print(json.dumps(inspect_cell(a.path, sh.strip("'"), ce), indent=1, default=str))
        return
    res = audit(a.path, recalc=not a.no_recalc, max_cells=a.max_cells)
    if a.json:
        json.dump(res, open(a.json, "w", encoding="utf-8"), indent=1)
    else:
        print(json.dumps(res, indent=1))
    sys.exit(1 if res["counts"].get("HIGH") else 0)


if __name__ == "__main__":
    main()
